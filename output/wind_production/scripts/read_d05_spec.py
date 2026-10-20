from docx import Document
p=r'C:\Code\UITxVNG\SE116_Nhom5_UITxVNG\output\documents\Map_Gio_Concept_va_Quy_Cach_Model.docx'
d=Document(p)
for i in range(145,174):
 if i<len(d.paragraphs):print(i,repr(d.paragraphs[i].text))
for ri,row in enumerate(d.tables[2].rows):
 print('ROW',ri,' | '.join(c.text.replace('\n',' / ') for c in row.cells))
