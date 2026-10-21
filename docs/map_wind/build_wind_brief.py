from pathlib import Path
import math
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

OUT=Path(__file__).parent
OUT.mkdir(parents=True,exist_ok=True)
CONCEPT=Path('C:/Users/DANG QUANG/.codex/generated_images/01a154c4-5257-7321-a108-206b165a7884/exec-6b0b01ed-4961-4bf3-9b03-d4dc56540a9d.png')

# Bản vẽ quy cách mới, không lấy hình học từ model cờ thử nghiệm.
im=Image.new('RGB',(1500,850),'#f6f5f1'); d=ImageDraw.Draw(im)
font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',25)
small=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',21)
def label(x,y,s):d.text((x,y),s,font=font,fill='#30343a')
def line(a,b,fill='#53585e',w=3):d.line([a,b],fill=fill,width=w)
def emblem(cx,cy,r):
    d.arc((cx-r,cy-r,cx+r,cy+r),35,325,fill='#4d565c',width=13)
    for dy in [-26,0,26]:
        d.arc((cx-r*.60,cy+dy-24,cx+r*.55,cy+dy+24),190,350,fill='#4d565c',width=10)
def flag(x,broken=False):
    line((x,745),(x,125),'#524c43',14)
    line((x-20,180),(x+178,180),'#524c43',10)
    pts=[(x+16,190),(x+166,190),(x+166,500),(x+92,475),(x+16,500)]
    if broken:pts=[(x+16,190),(x+166,190),(x+166,415),(x+140,442),(x+148,465),(x+116,448),(x+92,475),(x+16,490)]
    d.polygon(pts,fill='#e8e7df',outline='#747b7c',width=3)
    emblem(x+91,315,53)
    d.rectangle((x-22,725,x+22,750),fill='#777a78')
flag(230);flag(710,True)
label(190,45,'F01  CỜ CHÍNH');label(650,45,'F02  CỜ RÁCH')
line((170,125),(170,745));line((158,125),(183,125));line((158,745),(183,745));label(65,420,'4,8 m')
line((420,190),(420,500));line((408,190),(432,190));line((408,500),(432,500));label(435,330,'2,4 m')
line((246,545),(396,545));line((246,535),(246,555));line((396,535),(396,555));label(275,560,'1,2 m')
# Silhouette tỷ lệ giản lược, không phải model nhân vật thay thế.
d.polygon([(1085,745),(1090,650),(1065,620),(1072,554),(1084,524),(1097,552),(1118,552),(1134,524),(1147,560),(1144,620),(1127,650),(1140,745),(1120,745),(1108,670),(1102,745)],fill='#bc713b')
label(1010,780,'CÁO  ~1,755 m')
label(1030,45,'MỐC TỶ LỆ')
d.text((50,805),'Thiết kế đề xuất • Mặt trước • Cùng tỷ lệ • Nét biểu tượng lớn, không có chữ',font=small,fill='#53585e')
im.save(OUT/'flag_spec.png')

doc=Document(); s=doc.sections[0]
s.page_width=Cm(21);s.page_height=Cm(29.7)
s.top_margin=Cm(1.7);s.bottom_margin=Cm(1.7);s.left_margin=Cm(1.8);s.right_margin=Cm(1.8)
for name in ['Normal','Title','Heading 1','Heading 2','Heading 3','Caption']:
 st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
 st.font.size=Pt(10.5 if name=='Normal' else 22 if name=='Title' else 16 if name=='Heading 1' else 11.5 if name=='Heading 2' else 9)
 st.paragraph_format.space_after=Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing=1.06
doc.styles['Heading 2'].paragraph_format.space_before=Pt(9)
doc.core_properties.title='Concept và quy cách model Map Gió'
doc.core_properties.subject='Thiết kế model theo tỷ lệ cáo và camera hiện tại'
doc.core_properties.author='Nhóm phát triển SE116'
footer=s.footer.paragraphs[0];footer.alignment=2
r=footer.add_run(); fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)
def p(t,style=None):return doc.add_paragraph(t,style)
def h(t):doc.add_heading(t,1)
def page(t):doc.add_page_break();h(t)
def pic(path,width=17.3):
 run=doc.add_paragraph().add_run();sh=run.add_picture(str(path),width=Cm(width));sh._inline.docPr.set('descr',Path(path).stem)
def table(headers,rows,widths):
 t=doc.add_table(rows=1, cols=len(headers));t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Cm(w)
 for c,txt in zip(t.rows[0].cells,headers):c.text=txt
 rep=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rep)
 for row in rows:
  for c,txt in zip(t.add_row().cells,row):c.text=str(txt)
 for row in t.rows:
  for i,c in enumerate(row.cells):
   c.width=Cm(widths[i]);pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');b.append(el)
   pr.append(b)
   mar=OxmlElement('w:tcMar')
   for side in ['top','left','bottom','right']:
    el=OxmlElement('w:'+side);el.set(qn('w:w'),'85');el.set(qn('w:type'),'dxa');mar.append(el)
   pr.append(mar)
   for pp in c.paragraphs:
    pp.paragraph_format.space_after=Pt(3)
    for rr in pp.runs:rr.font.size=Pt(9.5)
 for c in t.rows[0].cells:
  for rr in c.paragraphs[0].runs:rr.bold=True
 return t
def asset(code,name,size,shape,damage,tech):
 doc.add_heading(code+'  '+name,2)
 p('Kích thước đề xuất: '+size+'.')
 p('Tạo hình: '+shape)
 p('Biến thể và lắp ghép: '+damage+' '+tech)

doc.add_heading('Concept và quy cách model Map Gió',0)
p('Tài liệu giao việc cho AI vẽ concept và dựng model 3D', 'Subtitle')
p('Thiết kế một nền văn minh từng vận hành bằng gió, nay đã sụp đổ. Mỗi model cần có dáng đẹp khi đứng riêng, ghép thành phế tích có trật tự và đọc rõ trong camera chơi. Lá cờ là dấu hiệu nhận diện bắt buộc của khu Gió.')
pic(CONCEPT)
p('Ảnh concept tổng thể đã chốt. Dùng để giữ ý tưởng kiến trúc, di tích và sự sụp đổ; giản lược chi tiết bề mặt khi dựng asset.', 'Caption')
doc.add_heading('Hướng tạo hình đã thống nhất',2)
p('Tham chiếu cách tổ chức hình khối và không gian của Death’s Door: dáng rõ, tỷ lệ có chủ ý, mảng vật liệu gọn và khoảng trống giúp nhân vật nổi bật. Bản thiết kế này áp dụng hướng đó cho kiến trúc Gió mang cảm hứng Đông Á; không sao chép asset của game tham chiếu.')
p('Ưu tiên: lá cờ → bộ khối kiến trúc đẹp → dấu vết sụp đổ có nguyên nhân → chi tiết phụ. Không dùng số lượng gạch vụn hoặc hoa văn để thay thế chất lượng tạo hình.')
p('Ngày đối chiếu project: 19 tháng 10 năm 2026. Phạm vi: môi trường Map Gió, không gồm vũ khí hoặc thay đổi cơ chế chơi.')

page('Cơ sở tỷ lệ từ cáo và camera')
p('Các số đo trong bảng này lấy từ nhân vật và cấu hình camera hiện tại. Kích thước model môi trường ở những trang sau là thiết kế mới dựa trên các mốc này, không kế thừa kích thước asset môi trường có sẵn.')
table(['Mốc kiểm tra','Giá trị','Ý nghĩa khi dựng'],[
['Chiều cao cáo','Khoảng 1,755 m','Từ y = 0 đến 1,75496 của mesh ở tư thế gốc trong GLB; chưa tính biến dạng animation.'],
['Capsule thân','Cao 1,8 m; bán kính 0,35 m','Dùng H = 1,8 m làm thước thiết kế; đường kính va chạm là 0,7 m.'],
['Camera tương đối','(0; 24,51; 20,57) m','Cách gốc nhân vật khoảng 32 m, góc nhìn xuống khoảng 50 độ.'],
['Ống kính và viewport','FOV 30 độ; 1280 × 720','Scene không đặt projection trực giao; đây là cấu hình phối cảnh hiện tại.'],
['Cáo trên màn hình','Ước tính khoảng 49 px chiều cao','Xấp xỉ 6,9% ở tâm ảnh, đứng thẳng trên mặt phẳng. Đây là tính hình học, chưa phải ảnh chụp runtime.'],
['Chi tiết đọc được','Đề xuất nét chính từ 0,12–0,20 m','Khoảng 5–9 px theo phương ngang gần nhân vật; mặt nghiêng có thể nhỏ hơn.']
],[4.0,5.0,8.3])
doc.add_heading('Quy tắc dựng cho AI',2)
p('Đơn vị mét; Y hướng lên. Trong tài liệu, R × S × C là rộng theo X × sâu theo Z × cao theo Y. Kích thước cờ ghi riêng rộng × cao của mặt vải. Pivot công trình đặt tại tâm đáy, mặt tiếp đất y = 0; module mặt sàn đặt pivot tại tâm mặt đi được, phần dày kéo xuống dưới y = 0.')
p('Không dùng bề rộng cáo trong tư thế dang tay làm bề rộng đi lại. Đuôi phải có chỗ hiển thị dù capsule hẹp hơn. Lối chính đề xuất thông thủy 4,5 m; cửa thông thủy 4,5 × 4,2 m; khoảng trống sân chính tối thiểu 12 × 12 m. Đây là mốc bố cục thử, không phải kết quả thử di chuyển.')
p('Giữ đúng kích thước ở cạnh nối module. Dung sai dựng đề xuất ±0,01 m tại đầu nối; phần trang trí được biến đổi trong bao ngoài đã ghi. Không thay khoảng nhảy, cao độ sàn hoặc collider của level chỉ để vừa một model mới.')
doc.add_heading('Cách kiểm tra tỷ lệ',2)
p('Đặt một cáo hiện tại cạnh model, dùng đúng camera trên và chụp 1280 × 720. Kiểm tra cả đứng gần và chạy ngang trước model. Mái, cột, cờ và mảnh vỡ phải giữ được dáng mà không che đầu, chân hay vị trí tiếp đất của cáo.')

page('Lá cờ nhận diện khu Gió')
p('Hạng mục bắt buộc và ưu tiên P0. Cờ được thiết kế mới từ tỷ lệ cáo và camera; không dùng cờ thử nghiệm làm mẫu. Mẫu dưới là bản vẽ quy cách đề xuất cho lần dựng đầu tiên.')
pic(OUT/'flag_spec.png')
table(['Mã','Quy cách','Cách tạo hình'],[
['F01 Cờ chính','Cột cao 4,8 m; mặt vải 1,2 × 2,4 m','Mép trên y = 4,3 m; đáy thấp nhất y = 1,9 m; đuôi chẻ chữ V nông 0,20 m. Thanh treo ngang dài 1,55 m.'],
['F02 Cờ rách','Giữ kích thước và điểm treo F01','Mất 15–25% phần đuôi hoặc mép ngoài; giữ nguyên biểu tượng và phần treo. Chỉ 2–3 vết xé lớn.'],
['F03 Cột gãy','Bao nằm ngang khoảng 4,8 × 1,5 × 0,8 m','Tách chân cột, đoạn thân đổ và vải vướng vào thân. Không dùng làm vật tương tác.'],
['F04 Cờ ngang phụ','Mặt vải 2,4 × 1,2 m; cột 4,2 m','Mép dưới y = 2,5 m; treo bên đường gió. Cùng biểu tượng; dùng thưa hơn cờ chính.']
],[3.2,5.6,8.5])
p('Biểu tượng: vòng hở ôm ba nét gió cong cùng chiều; đường kính 0,85 m, nét khoảng 0,12 m. Đặt ở nửa trên cờ; không viết chữ nhỏ. Màu đề xuất: vải #E8E7DF, dấu #4D565C, cột #524C43. Đây là bảng màu thiết kế, không phải màu bắt buộc đã có trong code.')
p('Cột thuôn nhẹ, tiết diện bát giác; chân đá 0,55 × 0,55 × 0,25 m; một đai kim loại tối ở đầu và chân. Vải có 2–3 nếp cong rộng khi đứng yên, không dựng thành tấm phẳng cứng. Tách vải khỏi cột; chừa vùng chuyển động phía sau 0,8 m làm mốc kiểm thử.')
p('Bố trí khởi điểm: 2 cờ cổng, 3 cờ hành lang, 2 cờ giếng, 1 cờ sân, 2 cờ bến, 1 cờ thư điện, 2 cờ đàn = 13 cờ đứng; thêm 2 cột gãy ngoài lối đi. Điều chỉnh theo từng khung hình; không rải đều như hàng rào.')

page('Bộ địa hình và mặt đi')
p('Tất cả kích thước dưới đây là đề xuất. P0 dựng thử, P1 bộ lõi, P2 hoàn thiện. Mỗi model dùng 1 khối chính, tối đa 2 khối phụ và một điểm nhấn rõ.')
asset('T01 • P0','Vách đá thẳng','4 × 3 × 6 m','Dáng trụ thuôn, chân rộng hơn đỉnh 10%; 3–5 mặt đá lớn lệch nhau, một đường nứt xiên sâu tạo nhịp. Đỉnh có mặt đặt sàn rõ.','Bản nguyên, khuyết góc và nứt; giữ cùng đầu nối rộng 4 m.','Pivot tâm đáy; phần chân được chìm vào đảo; ≤ 5.000 tam giác.')
asset('T02 • P1','Góc vách đá','4 × 4 × 6 m','Góc ngoài 90 độ với 2 mặt chính; vát một cạnh lớn để tránh cảm giác hộp vuông.','Bản lồi, lõm và gãy; khớp chiều cao T01.','Chừa mép nối phẳng; không thêm gờ có vẻ đi được ở dưới đường chơi.')
asset('T03 • P0','Tấm sàn kiến trúc','4 × 4 × 0,30 m','Một tấm lớn, chia tối đa 4 phiến; viền dày 0,15 m và một góc sứt. Mặt đi phẳng, cạnh ngoài có bevel hẹp.','Nguyên, nứt kín, sứt mép; vết nứt kín không thành hố thật.','Pivot mặt đi; đầu nối ±2 m; collider phẳng đơn giản.')
asset('T04 • P1','Mép sàn bị sạt','4 × 2 × 1 m','Mặt trên còn một dải nguyên, phần dưới vỡ thành 2–3 tầng lớn; vết vỡ khác rõ mặt xây thẳng.','Trái, phải và giữa; mảnh rơi phải có mặt gãy tương ứng.','Đặt ngoài phần đường cần bảo toàn; không lấn khoảng đáp.')
asset('T05 • P1','Cầu thang nghi lễ','4,5 × 3,6 × 1,8 m','9 bậc cao 0,20 m, sâu 0,40 m; hai má đá thấp, một đầu má sứt. Nhịp bậc đều và có chiếu nghỉ riêng.','Nguyên và hỏng má bên; không phá mặt bậc bắt buộc.','Pivot chân thang; thử collider dốc liên tục thay vì cho capsule va từng cạnh bậc.')

page('Bộ kiến trúc thanh thoát')
asset('A01 • P0','Cột nghi lễ','0,8 × 0,8 × 4,2 m','Chân 0,8 m, thân bát giác rộng 0,45 m thuôn nhẹ, đầu cột 0,7 m. Nhịp chân–thân–đầu rõ; một đai ngang là điểm nhấn.','Nguyên, gãy còn 1,7 m và đầu cột rơi riêng.','Vết gãy xiên 20–30 độ, 3 mặt lớn; không dùng răng cưa nhỏ.')
asset('A02 • P1','Dầm ngang','4 × 0,45 × 0,6 m','Thân dầm gọn; hai đầu nhô nhẹ, đường dưới cong rất nông. Dùng tiết diện rõ thay vì khắc vân dày.','Nguyên và gãy đôi với hai mặt khớp.','Pivot giữa đáy; dựng theo X; điểm nối ±2 m.')
asset('A03 • P1','Tường thấp','4 × 0,35 × 1 m','Thân phẳng, chân và nắp có tỷ lệ rõ; nắp dày 0,12 m. Một mảng khuyết lớn tạo nhịp.','Nguyên, gãy thấp và đầu kết thúc.','Phía gần camera ưu tiên thấp; không xây thành tường che cáo.')
asset('A04 • P1','Mái điện','6 × 4 × 1,4 m','Hai mặt mái lớn; sống mái gọn, bốn góc hất nhẹ. Chỉ 5–7 dải ngói lớn gợi nhịp, không dựng từng viên.','Nguyên và mất một góc 20%; để lộ 2–3 xà.','Mái tách khỏi cột; phần mái trước tháo riêng để xử lý che khuất.')
asset('A05 • P1','Lan can nghi lễ','4 × 0,25 × 0,9 m','Hai trụ đầu 0,3 m, một thanh ngang và một khoảng rỗng lớn; cùng ngôn ngữ A01.','Nguyên, mất thanh và trụ đổ.','Đầu nối đồng bộ sàn; tách phần trang trí khỏi rào gameplay.')

page('Cầu và hạ tầng dùng gió')
asset('I01 • P1','Nhịp cầu mới','4 × 4,5 × 0,45 m','Dài theo X; mặt đi rộng 4,5 m theo Z. Sàn gồm 3–4 mảng lớn; dầm dưới nhô vừa đủ, không nhiều dây nhỏ.','Nguyên và gãy chéo; bản gãy dùng khi layout cho phép hố.','Đầu nối X = ±2 m; ghép nhiều nhịp; không lấy kích thước cầu cũ.')
asset('I02 • P1','Trụ đỡ cầu','1,2 × 1,2 × 5,4 m','Chân vững, thân thuôn 15%, đầu chạc ôm dầm; một đường phân lớp đá.','Nguyên và nghiêng gãy có mảnh chân tương ứng.','Tách đoạn thân để tăng chiều cao bằng module, không kéo giãn toàn bộ.')
asset('I03 • P0','Máy buồm đón gió','4,8 × 2 × 6,3 m','Tháp chân 1,5 m; rotor đường kính 4,5 m, tâm cao 4 m; 4 cánh buồm dạng hình thang, giữa có trục lớn rõ.','Hoạt động đủ cánh; hỏng thiếu 1 cánh và rách một cánh; bản đổ ngoài đường.','Tách Rotor, Frame, Sail; pivot rotor đúng tâm trục; collider phần cố định đơn giản.')
asset('I04 • P1','Trụ dẫn gió','0,9 × 0,9 × 3,6 m','Thân đá hẹp, đỉnh mở một khe cong lớn hướng theo luồng gió; cùng biểu tượng với cờ nhưng tiết chế.','Nguyên và vỡ phần đỉnh.','Trạng thái hoạt động do game điều khiển; không bake ánh sáng phát quang vào bản hỏng.')
asset('I05 • P1','Bến gỗ','4,5 × 4 × 0,6 m','5–7 tấm ván rộng; 4 cọc chân tạo nhịp; mép thuyền cập phẳng. Mũi bến vát nhẹ thành dáng riêng.','Nguyên, ván hỏng ngoài mép và góc sụp; giữ diện tích đứng.','Tách cọc chân để lắp theo mực nước thực tế; không đổi điểm cập thuyền.')

page('Thuyền và các điểm sử dụng')
p('Các mẫu dưới đây cũng được thiết kế mới theo H = 1,8 m. Khi tích hợp, cần kiểm tra riêng tuyến thuyền, điểm cập và vùng tương tác; không mặc định hình học mới sẽ khớp logic hiện tại.')
asset('I06 • P1','Thuyền buồm Gió','2,7 × 5,4 × 4,8 m','Thân thuyền thuôn với mũi hơi nâng; boong phẳng dài 4 m, rộng giữa 2,2 m. Một cột buồm và một mảng buồm tam giác lớn, không nhiều dây chằng. Dấu Gió đặt trên phần rộng của buồm.','Thuyền dùng được giữ đủ buồm; xác thuyền trang trí lật nghiêng, mất cột.','Tách Hull, Mast, Sail; mặt boong và điểm neo phải ghi trong file nguồn. Không sửa tuyến chạy ở bước dựng model.')
asset('I07 • P2','Cọc neo thuyền','0,45 × 0,45 × 1,2 m','Cọc gỗ bát giác, đầu loe 0,55 m, một vòng dây lớn quấn thấp. Chân chìm vào mép bến.','Đứng nguyên và nghiêng gãy; không thêm nhiều cuộn dây nhỏ.','Tách dây nếu cần chuyển động; đặt ngoài vùng cáo bước lên boong.')
asset('I08 • P2','Cần nâng hàng bỏ hoang','3,6 × 1,8 × 4,5 m','Một cột đứng, tay đòn chéo dài và bánh ròng rọc lớn 0,7 m. Silhouette tam giác rõ, giỏ hàng gọn phía dưới.','Bản tay đòn gãy và giỏ hàng rơi cùng hướng.','Không tạo vẻ như câu đố đang dùng nếu chỉ là trang trí; đặt ở mép bến xa camera.')
asset('P03 • P1','Bệ sách Gió','1,5 × 1,2 × 1,5 m','Bệ đá thuôn, mặt đọc nghiêng nhẹ và một cuốn sách lớn 0,7 × 0,5 m. Một đường cong gợi luồng gió dẫn từ chân lên mặt bệ.','Bệ có góc sứt; sách còn nguyên là tâm chú ý.','Tách sách khỏi bệ; để trống ít nhất 1,5 m phía tiếp cận cho tương tác.')
asset('P04 • P1','Điểm lưu phong cách Gió','Đế đường kính 3,6 m, cao 0,18 m; tinh thể cao 0,9 m','Vòng thấp chia ba khối sạch, tinh thể ở giữa có 5–6 mặt lớn. Tâm tinh thể cao 1,8 m; nền tối, ánh sáng ấm khi kích hoạt.','Trạng thái tắt và sáng; không dùng bản đổ ở điểm lưu đang hoạt động.','Tách đế và tinh thể; phần sáng do gameplay điều khiển. Đế thấp không làm thành vật cản trên đường.')

page('Công trình điểm nhấn')
p('P2. Đây là bao kích thước mục tiêu của cụm, không phải một mesh khổng lồ. Dùng lại bộ module đã duyệt; giữ mặt đi và khoảng trống quan trọng của level.')
asset('L01','Cổng phế tích','7,2 × 2,4 × 5,8 m','Hai chân cột lớn, mái một lớp gọn và thanh ngang; cửa trống rộng 4,5 m, cao ít nhất 4,2 m. Hai cờ chính đặt lệch ra ngoài.','Một góc mái mất; một cột sứt nhưng cổng còn đọc rõ dáng.','Chia cột, dầm, mái, chân; không đặt đống đá ở giữa cửa.')
asset('L02','Giếng Gió','Đường kính ngoài 12 m; trong 8 m; vành cao 1,2 m','Vòng đá ngang chia 8 cung, viền trên rộng rõ; hai vị trí tiếp cận mở theo đường chơi. Có 3 trụ thấp tạo nhịp.','Một cung ngoài nứt trượt; đường dẫn gió còn liên tục.','Chừa lõi rỗng cho vùng gió; chiều cao luồng do gameplay quyết định; không bịt bằng mesh.')
asset('L03','Thư điện Gió','7,2 × 6 × 5,8 m','Bệ thấp, bốn cột, mái A04 và gian giữa thông thoáng. Đặt bệ sách đơn giản ở hậu cảnh; nhìn thấy lối vào từ camera.','Mất góc mái và một mảng tường sau; bên trong vẫn có cấu trúc.','Không làm thư viện đầy sách nhỏ; giữ khoảng tương tác quanh vật phẩm.')
asset('L04','Đàn Gió và vòng vỡ','Sân 16 × 16 m; vòng ngoài 10 m, trong 7 m, dày 0,9 m','Vòng đứng trên hai chân thấp, một rãnh lớn đi theo cung. Vòng nằm ở phía xa camera, phía trước là sân trống tối thiểu 12 × 12 m.','Mất cung 70–90 độ; mảnh cùng bán kính nằm ở vách bên dưới.','Vòng tách 8–12 cung; không mặc định đây là cổng tương tác hoặc boss đã có.')
asset('L05','Tháp quan trắc gãy','3 × 3 × 9 m','Chân vuông vững, thân tháp ba tầng khối lớn, đỉnh từng có khung buồm. Dáng vươn cao, khoảng rỗng rõ.','Bản còn 6 m và đoạn đỉnh 3 m rơi cạnh chân.','Đặt ở hậu cảnh; tách tầng để tối ưu che khuất và lặp lại.')

page('Dấu vết đời sống và sụp đổ')
asset('D01 • P1','Mảnh vòng nghi lễ','Bao khoảng 5 × 2 × 4 m','Lấy chính cung bị mất của L04, giữ bán kính, tiết diện và rãnh. Xoay nghiêng rồi bổ sung 2–3 mặt vỡ lớn.','Một mảnh chính và tối đa hai mảnh phụ.','Kích thước bao phụ thuộc góc đặt; tuyệt đối không làm mảnh tròn không khớp vòng.')
asset('D02 • P1','Cụm đá đổ','Nhỏ 1,2 × 1 × 0,5 m; lớn 3 × 2 × 1,2 m','Một khối chủ đạo và 3–5 khối phụ; khoảng hở tạo bóng. Mặt xây còn vuông, mặt gãy bất quy tắc.','Ba bố cục khác nhau, không xoay cùng một cụm khắp map.','Đá sát đúng công trình bị hỏng; tránh rải thành nhiễu trên lối đi.')
asset('D03 • P2','Xe hàng bỏ lại','1,4 × 2,2 × 1,3 m','Thùng gỗ hẹp, hai bánh đĩa rõ; tay kéo tạo đường chéo dài. Chỉ hai nẹp kim loại lớn.','Đứng nguyên, lật mất bánh; hàng rơi về một phía.','Đặt ở bến hoặc kho, gợi công việc bị gián đoạn; không mặc định có thể đẩy.')
asset('D04 • P2','Hòm và giá hàng','Hòm 0,9 × 0,7 × 0,65 m; giá 1,8 × 0,7 × 1,6 m','Hòm có nắp rõ và một đai; giá có 2 tầng thưa. Tạo hình gọn, màu gỗ trầm, không chữ nhỏ.','Hòm mở rỗng, giá gãy một chân; 2–3 hòm là đủ cho một cụm.','Dùng chung vật liệu; đặt lệch tuyến chạy.')
asset('D05 • P2','Phù điêu Gió','2,4 × 0,35 × 2,8 m','Phiến đá đứng có vòng hở và ba nét gió lớn đồng nhất với cờ; relief sâu khoảng 0,10 m.','Nguyên và mất góc dưới; không làm mất hết biểu tượng.','Nền tối hơn mặt nổi một bậc; không dùng chữ khắc li ti.')

page('Thiên nhiên và vật phụ trợ')
asset('N01 • P1','Cây nghiêng theo gió','Tán 4,5 × 3 m; cao 4,8 m','Thân cong một hướng, 2–3 nhánh chính; tán gồm 4–6 khối lá phẳng, đầu tán kéo theo gió.','Một cây sống thưa và một cây khô còn 2 nhánh lớn.','Tránh tán gần camera che đường; chi tiết lá không cần từng chiếc.')
asset('N02 • P2','Bụi cây và cỏ','Bụi 1,2 × 1 × 0,65 m; cỏ 0,7 × 0,5 × 0,4 m','Cụm thấp lệch về hướng gió, ba mảng chính; giá trị sáng tối gần nhau.','3 silhouette; không đổi màu quá nhiều.','Dùng ở chân tường và kẽ đá; mật độ thấp trên đường chơi.')
asset('N03 • P2','Rễ và dây leo','Dài 2,4 m; bao rộng 0,6 m; nhô 0,12 m','Một đường chính cùng 2 nhánh lớn, ôm cạnh đá hoặc thả xuống vách.','Thẳng, rẽ góc và rủ; không phủ mọi mặt công trình.','Ưu tiên silhouette ở mép, bỏ chi tiết không thấy được từ camera.')
asset('P01 • P2','Đèn nghi lễ','0,5 × 0,5 × 1,6 m','Bệ gọn, thân thanh, lồng đèn có một ô sáng lớn. Mái đèn chỉ hai mặt dốc.','Tắt, vỡ và bản sáng ở vị trí có chủ đích.','Không để đèn trang trí cạnh tranh màu checkpoint; phần sáng tách material.')
asset('P02 • P1','Tường giữ hồ','4 × 1 × 2,4 m','Mặt kè lớn có chân loe; nắp rộng tạo đường viền hồ. Khe xả nước là một khoảng rỗng lớn.','Nguyên, góc 90 độ, vỡ phía ngoài.','Phần giữ hồ còn liên tục; không tạo cảnh nước treo vô lý hoặc thông thẳng xuống biển mây.')

page('Bố cục và trình tự sản xuất')
table(['Cụm khu vực','Vai trò hình ảnh','Model ưu tiên'],[
['Cổng và đường đầu','Nhận ra khu Gió ngay; dấu hư hỏng nhẹ','L01, F01, T01–T05, A01'],
['Hành lang đón gió','Cờ và trụ cho thấy hướng gió; máy còn hoạt động','F04, I03, I04, A03'],
['Giếng và cầu','Nhìn ra độ cao; thấy hạ tầng đã gãy','L02, I01–I02, F02, D02'],
['Sân hộ vệ','Khoảng trống chiến đấu; viền kiến trúc tạo khung','T03, A01–A05, D05'],
['Bến và thư điện','Dấu đời sống bỏ dở; liên hệ thuyền với gió','I05, L03, D03–D04, F01'],
['Hồ và bờ xa','Mặt nước thoáng; tầm nhìn dẫn đến vòng vỡ','P02, L04, D01, L05']
],[3.5,7.3,6.5])
doc.add_heading('Ba tầng dấu hiệu sụp đổ',2)
p('Toàn cảnh: tháp gãy, vòng mất cung, cầu đứt. Cự ly vừa: rotor thiếu cánh, mái sập, trụ nghiêng. Cự ly gần: cờ rách, xe hàng lật, phù điêu vỡ. Các mảnh rơi phải khớp nơi chúng tách ra; giữ nguyên nhân thảm họa ở mức bí ẩn.')
doc.add_heading('Thứ tự dựng và điều kiện chuyển bước',2)
p('P0: dựng F01, A01, T01, T03 và I03; đặt cạnh cáo bằng camera thật. Chỉ mở rộng khi từng model đẹp ở hình đen, vật liệu xám và ánh sáng trong game.')
p('P1: hoàn thiện module đường đi và bản hư hỏng; ghép một cụm Giếng Gió có cầu, cờ và vách. Đi thử để kiểm tra tầm nhìn, mép sàn và đường đáp. Sau đó mới phủ toàn map.')
p('P2: xây các công trình điểm nhấn, cây và dấu đời sống. So từng cụm từ camera, bỏ chi tiết không đọc được. Đặt mảnh vỡ theo mối liên hệ cấu trúc, không theo mật độ trang trí.')
p('Mốc nghiệm thu đầu tiên: một cảnh mẫu chứa cáo, cờ chính, cột nguyên và gãy, sàn, vách, máy buồm. Mẫu này quyết định chất lượng của toàn bộ bộ model; chưa cần làm cả map.')

page('Yêu cầu giao cho AI và nghiệm thu')
doc.add_heading('Prompt nền dùng cho mọi model',2)
p('Thiết kế asset 3D stylized cho game hành động nhìn chéo từ trên cao. Hướng thẩm mỹ tham chiếu Death’s Door: hình khối tinh giản nhưng được trau chuốt, tỷ lệ thanh thoát, silhouette rõ, ít vật liệu và khoảng trống có chủ ý. Chủ đề nền văn minh Gió Đông Á đã sụp đổ, đá xám lạnh, gỗ nâu tối, cờ trắng ngà. Cáo đứng hai chân cao 1,755 m, chuẩn gameplay H = 1,8 m. Camera phối cảnh nhìn xuống 50 độ, FOV 30 độ, cách nhân vật khoảng 32 m; viewport 1280 × 720. Dựng đúng quy cách của mã model đi kèm. Mỗi asset có một dáng chính đẹp, một điểm nhấn rõ, biến thể hỏng có cấu trúc. Tránh đầu to kiểu đồ chơi, bề mặt siêu thực, vân đá dày, hàng trăm viên gạch, hoa văn nhỏ và ánh sáng phát quang trang trí khắp nơi.')
doc.add_heading('Đầu ra cho từng mã model',2)
p('Giai đoạn vẽ: một bảng nền trung tính có mặt trước, bên, trên và phối cảnh ba phần tư; cùng tỷ lệ giữa các hình; có cáo làm thước và kích thước ghi rõ. Xuất thêm một ảnh bằng góc camera game. Bản vẽ kích thước là chỉ dẫn; AI tạo ảnh không bảo đảm hình học chính xác, phải đo lại khi dựng 3D.')
p('Giai đoạn model: GLB theo mét, Y lên; file nguồn có các bộ phận đặt tên; apply scale và kiểm tra bao kích thước. Một vật liệu chung mỗi model theo art guide; phần đổi trạng thái được thêm một vật liệu riêng. Khối môi trường ≤ 5.000 tam giác mỗi module; vật tương tác vừa ≤ 2.000. Cụm lớn chia module, không ép tất cả chi tiết vào một mesh.')
p('Tách chuyển động: vải, rotor và cánh buồm; pivot rotor ở trục; UV vải đầu cố định u = 0, mép tự do u = 1. Cờ dọc cần cách neo theo thanh ngang ở mép trên; không giả định shader thử nghiệm hiện tại áp dụng trực tiếp. Thiết kế biểu tượng cần material hoặc texture hỗ trợ riêng khi tích hợp.')
doc.add_heading('Tiêu chí đẹp và đúng trong game',2)
p('1. Nhìn bóng đen vẫn nhận ra model. 2. Dùng vật liệu xám vẫn thấy tỷ lệ đẹp. 3. Trong 720p đọc được điểm nhấn. 4. Hai bản nguyên và gãy vẫn cùng một thiết kế. 5. Cáo và đường đi không bị che. 6. Cờ đọc được biểu tượng ở vùng nhìn chính. 7. Kích thước và điểm nối đúng bảng. 8. Không đổi collider hoặc đường đi ngoài phạm vi tích hợp được duyệt.')
doc.add_heading('Nguồn đối chiếu',2)
p('Project: entities/player/player.tscn; entities/_base/character.tscn; entities/player/camera_rig.gd; generated/models/fox_spirit.glb; generated/fox_player.gd; project.godot. Phong cách và ngân sách: docs/art_style_guide.md. Cấu trúc khu vực: levels/zone_1_wind/level_1/level_1.tscn.')
p('Tham chiếu hình ảnh Death’s Door: trang chính thức https://www.playdeathsdoor.com/ và nhà phát hành https://www.devolverdigital.com/games/deaths-door. Các quy cách tạo hình trong tài liệu là định hướng thiết kế cho project này, không phải thông số kỹ thuật của Death’s Door.')

dest=OUT/'Map_Gio_Concept_va_Quy_Cach_Model.docx'
doc.save(dest)
print(dest.resolve())
print('Pages planned: 12; asset specifications: 39 (4 flags + 35 environment assets)')
